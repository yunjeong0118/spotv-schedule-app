# -*- coding: utf-8 -*-
"""
채널별 누적 업로드 마스터 + 주차별/기준시각별 업로드파일 생성 엔진 (AI 미사용, 규칙 기반)

- 마스터: 채널당 1개 xlsx. 편성 1건 = 1행 (절대 일시 기준). 주차와 무관하게 계속 쌓인다.
- 편성표 적용: 편성표가 다루는 구간(주 시작 ~ +7일)의 기존 데이터를 통째로 교체한다
  (변경 편성표 재적용 = 개정). 구간 밖 데이터는 절대 건드리지 않는다.
- 주차 파일: 월요일 + 기준시각(03:00 / 00:00)부터 +7일까지를 마스터에서 잘라낸다.
- 커버리지(적용 구간)를 기록해, 잘라낸 구간에 데이터가 없는 부분(누락)을 알려준다.
"""
import io
from datetime import datetime, date, timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill

from schedule_parser import (
    entry_dt, entry_key, entry_iyk, diff_entries, write_upload_excel, find_start_date,
    DAY_COLS_DEFAULT,
)

# ---------------------------------------------------------------------------
# 채널 설정: 업로드파일 작성 기준 시각(시)
# ---------------------------------------------------------------------------
CHANNELS = {
    'SPOTV':        {'bases': [3]},
    'SPOTV2':       {'bases': [3]},
    'SPOTV Golf+':  {'bases': [3]},
    'SPOTV Plus':   {'bases': [0, 3]},   # 두 버전 모두 작성
    'SPOTV Prime':  {'bases': [0]},
    'SPOTV Prime2': {'bases': [0]},
    'SPOTV Prime+': {'bases': [0]},
}

MASTER_SHEET = '마스터'
LOG_SHEET = '적용이력'
INFO_SHEET = '정보'
DT_FMT = '%Y-%m-%d %H:%M'


# ---------------------------------------------------------------------------
# 주차 계산
# ---------------------------------------------------------------------------
def monday_of(d):
    return d - timedelta(days=d.weekday())


def week_label(monday):
    """'N월 M주차' — 그 주의 목요일이 속한 달 기준 (8/17주=8월3주차, 9/7주=9월2주차)."""
    thu = monday + timedelta(days=3)
    nth = (thu.day - 1) // 7 + 1
    return thu.year, thu.month, nth


def week_filename(channel, monday, base_hour):
    y, m, n = week_label(monday)
    ch = channel.replace(' ', '_')
    return f"{y}년_{m}월_{n}주차_{ch}_업로드파일_{base_hour:02d}시기준_{monday:%y%m%d}.xlsx"


# ---------------------------------------------------------------------------
# 편성표가 다루는 구간 (적용 구간)
# ---------------------------------------------------------------------------
def grid_window(file_bytes, day_cols=None):
    """편성표 그리드의 시작/끝 일시. 시작 = 첫 날짜 열의 날짜 + B열 첫 시 라벨, 길이 = 날짜 열 수 × 1일."""
    day_cols = day_cols or DAY_COLS_DEFAULT
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    day_dates = find_start_date(ws, day_cols)
    if not day_dates:
        raise ValueError('편성표 헤더에서 날짜를 찾지 못했습니다.')
    first = min(day_dates.values())
    b = sorted((mc.min_row, ws.cell(row=mc.min_row, column=2).value)
               for mc in ws.merged_cells.ranges if mc.min_col == 2)
    first_hour = None
    for _, label in b:
        try:
            first_hour = int(label) % 24
            break
        except (TypeError, ValueError):
            continue
    if first_hour is None:
        raise ValueError('B열 시간 라벨을 찾지 못했습니다.')
    start = datetime.combine(first, datetime.min.time()) + timedelta(hours=first_hour)
    end = start + timedelta(days=len(day_dates))
    return start, end


# ---------------------------------------------------------------------------
# 마스터 I/O
# ---------------------------------------------------------------------------
def new_master(channel):
    return {'channel': channel, 'entries': [], 'log': []}


def load_master(file_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    channel = None
    if INFO_SHEET in wb.sheetnames:
        v = wb[INFO_SHEET].cell(row=1, column=2).value
        channel = str(v) if v else None
    entries = []
    ws = wb[MASTER_SHEET]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] is None:
            continue
        flags = set()
        if row[8] == 'Y': flags.add('자')
        if row[9] == 'Y': flags.add('해')
        if row[10] == 'Y': flags.add('수')
        entries.append({
            'date': date(int(row[1]), int(row[2]), int(row[3])),
            'hour': int(row[4]), 'minute': int(row[5]),
            'title': row[6], 'cat': row[7], 'flags': flags,
            'needs_review': (row[11] == 'Y') if len(row) > 11 else False,
            'src': (row[12] if len(row) > 12 else None),
        })
    log = []
    if LOG_SHEET in wb.sheetnames:
        for r in wb[LOG_SHEET].iter_rows(min_row=2, values_only=True):
            if not r[0]:
                continue
            log.append({
                'applied_at': r[0], 'source': r[1],
                'start': datetime.strptime(r[2], DT_FMT), 'end': datetime.strptime(r[3], DT_FMT),
                'added': r[4], 'removed': r[5], 'changed': r[6],
            })
    return {'channel': channel, 'entries': entries, 'log': log}


def write_master(master):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = MASTER_SHEET
    headers = [f'str{i}' for i in range(1, 12)] + ['검토필요', '출처']
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h).font = Font(name='Arial')
    ORANGE = PatternFill(start_color='FFD8A8', end_color='FFD8A8', fill_type='solid')
    for i, e in enumerate(sorted(master['entries'], key=entry_dt), 1):
        fl = e.get('flags', set())
        vals = [str(i), e['date'].year, f"{e['date'].month:02d}", f"{e['date'].day:02d}",
                e['hour'], e['minute'], e['title'], e['cat'],
                'Y' if '자' in fl else None, 'Y' if '해' in fl else None, 'Y' if '수' in fl else None,
                'Y' if e.get('needs_review') else None, e.get('src')]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i + 1, column=c, value=v)
            cell.font = Font(name='Arial')
            if c in (5, 6):
                cell.number_format = '00'
                if e.get('needs_review'):
                    cell.fill = ORANGE
    for c, w in {1: 6, 2: 6, 3: 5, 4: 5, 5: 5, 6: 5, 7: 55, 8: 9, 9: 6, 10: 6, 11: 6, 12: 8, 13: 40}.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = 'A2'

    lg = wb.create_sheet(LOG_SHEET)
    for c, h in enumerate(['적용일시', '원본 편성표', '적용구간 시작', '적용구간 끝', '추가', '삭제', '변경'], 1):
        lg.cell(row=1, column=c, value=h).font = Font(name='Arial')
    for i, g in enumerate(master['log'], 2):
        for c, v in enumerate([g['applied_at'], g['source'], g['start'].strftime(DT_FMT),
                               g['end'].strftime(DT_FMT), g['added'], g['removed'], g['changed']], 1):
            lg.cell(row=i, column=c, value=v).font = Font(name='Arial')
    for c, w in {1: 18, 2: 45, 3: 18, 4: 18}.items():
        lg.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w

    info = wb.create_sheet(INFO_SHEET)
    info.cell(row=1, column=1, value='채널')
    info.cell(row=1, column=2, value=master['channel'])
    info.cell(row=2, column=1, value='설명')
    info.cell(row=2, column=2, value='채널별 누적 업로드 마스터. 직접 편집하지 말고 앱에서 편성표를 적용하세요.')
    return wb


def master_bytes(master):
    bio = io.BytesIO()
    write_master(master).save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# 편성표 적용 (구간 교체)
# ---------------------------------------------------------------------------
def apply_schedule(master, new_entries, window, source, applied_at=None):
    """마스터의 window 구간 데이터를 new_entries로 교체한다. 구간 밖은 보존.
    반환: (report dict). master는 제자리에서 갱신된다.
    """
    applied_at = applied_at or datetime.now().strftime(DT_FMT)
    start, end = window
    notes = []

    dts = [entry_dt(e) for e in new_entries]
    if dts:
        if min(dts) < start:
            notes.append(f"그리드 시작({start:%m/%d %H:%M}) 이전 시각 항목 {min(dts):%m/%d %H:%M} 포함 → 적용구간 시작을 앞당김")
            start = min(dts)
        if max(dts) >= end:
            notes.append(f"그리드 끝({end:%m/%d %H:%M}) 이후 시각 항목 {max(dts):%m/%d %H:%M} 포함 → 적용구간 끝을 늘림")
            end = max(dts) + timedelta(minutes=1)

    old_in = [e for e in master['entries'] if start <= entry_dt(e) < end]
    keep = [e for e in master['entries'] if not (start <= entry_dt(e) < end)]

    # 같은 시각 중복 방지 (편성표 내부 중복)
    seen = {}
    for e in new_entries:
        k = entry_key(e)
        if k in seen:
            notes.append(f"편성표 내 동일 시각 중복: {k} — 마지막 항목 사용 (확인 필요)")
        seen[k] = e
    cleaned = []
    for e in seen.values():
        e2 = dict(e)
        e2['src'] = source
        cleaned.append(e2)

    changes = diff_entries(old_in, cleaned)
    cnt = {'ADDED': 0, 'REMOVED': 0, 'CHANGED': 0}
    for c in changes:
        cnt[c['kind']] += 1

    master['entries'] = keep + cleaned
    master['entries'].sort(key=entry_dt)
    master['log'].append({'applied_at': applied_at, 'source': source, 'start': start, 'end': end,
                          'added': cnt['ADDED'], 'removed': cnt['REMOVED'], 'changed': cnt['CHANGED']})
    return {'window': (start, end), 'old_count': len(old_in), 'new_count': len(cleaned),
            'changes': changes, 'counts': cnt, 'notes': notes, 'first_apply': len(old_in) == 0}


# ---------------------------------------------------------------------------
# 커버리지 / 주차 슬라이스
# ---------------------------------------------------------------------------
def coverage(master):
    """적용 구간들의 합집합 [(start,end), ...]"""
    iv = sorted((g['start'], g['end']) for g in master['log'])
    merged = []
    for s, e in iv:
        if merged and s <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], e))
        else:
            merged.append((s, e))
    return merged


def missing_intervals(master, start, end):
    """[start,end) 중 적용 구간에 포함되지 않은 부분"""
    gaps, cur = [], start
    for s, e in coverage(master):
        if e <= cur:
            continue
        if s >= end:
            break
        if s > cur:
            gaps.append((cur, min(s, end)))
        cur = max(cur, e)
        if cur >= end:
            break
    if cur < end:
        gaps.append((cur, end))
    return gaps


def slice_week(master, monday, base_hour):
    """월요일 + base_hour 부터 7일치 항목과 점검 정보를 반환."""
    start = datetime.combine(monday, datetime.min.time()) + timedelta(hours=base_hour)
    end = start + timedelta(days=7)
    entries = sorted([e for e in master['entries'] if start <= entry_dt(e) < end], key=entry_dt)
    gaps = missing_intervals(master, start, end)
    notes = []
    if entries and entry_dt(entries[0]) > start and not gaps:
        prev = [e for e in master['entries'] if entry_dt(e) < start]
        if prev:
            p = max(prev, key=entry_dt)
            notes.append(f"기준 시각({start:%m/%d %H:%M})에 진행 중인 편성 [{entry_dt(p):%m/%d %H:%M} 시작] {p['title']} 은(는) "
                         f"시작이 기준 시각 이전이라 이 파일에 포함되지 않음 (직전 주 파일에 포함). 첫 행은 {entry_dt(entries[0]):%m/%d %H:%M}")
        else:
            notes.append(f"첫 편성이 {entry_dt(entries[0]):%m/%d %H:%M}에 시작 — 기준 시각({start:%H:%M})에 걸친 직전 편성은 마스터에 없음. 확인 필요")
    review = [e for e in entries if e.get('needs_review')]
    return {'start': start, 'end': end, 'entries': entries, 'gaps': gaps,
            'notes': notes, 'review': review}


def fmt_gap(g):
    s, e = g
    return f"{s:%m/%d %H:%M} ~ {e:%m/%d %H:%M}"


def week_workbook(entries, channel, base_hour):
    return write_upload_excel(entries, None, sheet_title=f"{channel}_{base_hour:02d}시기준")


def week_bytes(entries, channel, base_hour):
    bio = io.BytesIO()
    week_workbook(entries, channel, base_hour).save(bio)
    return bio.getvalue()
